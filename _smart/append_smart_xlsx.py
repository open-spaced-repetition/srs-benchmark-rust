"""Append the KL and optimal-partition smart-preset configs to Smart Preset Assignment.xlsx.

The sheet already holds the Mahalanobis runs: baseline (row 3), 30 hierarchical (4-33), 16 HDBSCAN
(34-49). This appends, right after the Maha block, in order:
  - 30 KL hierarchical + 16 KL HDBSCAN   (column B = FSRS-7-short-secs-smart-kl)
  - 4 optimal partition (BIC/AIC x Maha/KL pre-merge)  (column B = ...-smart-opt)
then moves the "1000 users / N reviews" footer below. The Δ columns (F/H/J/L) reuse the existing
`=(x-$3)/$3` formulas vs the shared baseline in row 3. Idempotent: re-running overwrites the same
rows (footer text is reconstructed). Configs whose files aren't at 1000 users yet are written with
the labels/formulas but blank values.
"""
import copy
import json
import os

import openpyxl

RDIR = "result"
XLSX = "Smart Preset Assignment.xlsx"
BASE = "FSRS-7-short-secs"
MAHA_LAST_ROW = 49  # KL block starts right after

METHODS = ["single", "complete", "average", "centroid", "ward"]
KL_THRESHOLDS = ["0.001", "0.003", "0.008", "0.02", "0.05", "0.15"]
HDB_MCS = [2, 5, 10, 20]
HDB_MS = [1, 5]


def load(name):
    p = f"{RDIR}/{name}.jsonl"
    if not os.path.exists(p):
        return None
    d = {}
    for line in open(p):
        line = line.strip()
        if not line:
            continue
        o = json.loads(line)
        if "metrics" in o and o["user"] <= 1000:
            d[o["user"]] = o
    return d or None


def metrics(d):
    us = sorted(d)
    ll = [d[u]["metrics"]["LogLoss"] for u in us]
    rb = [d[u]["metrics"]["RMSE(bins)"] for u in us]
    sz = [d[u]["size"] for u in us]
    n, tot = len(us), sum(sz)
    return (sum(ll) / n, sum(rb) / n,
            sum(v * s for v, s in zip(ll, sz)) / tot,
            sum(v * s for v, s in zip(rb, sz)) / tot)


# (file, column-B version, C label, D method, style-template) in sheet order.
configs = []
for m in METHODS:
    for t in KL_THRESHOLDS:
        configs.append((f"{BASE}-smart-kl-{m}-{t}", f"{BASE}-smart-kl", float(t), m, "hier"))
for mcs in HDB_MCS:
    for ms in HDB_MS:
        for v in ("eom", "leaf"):
            configs.append((f"{BASE}-smart-kl-hdbscan-mcs{mcs}-ms{ms}-{v}", f"{BASE}-smart-kl",
                            f"mcs={mcs}, ms={ms}", f"HDBSCAN ({v})", "hdb"))
# Optimal in sweep order: bic-maha, aic-maha, bic-kl, aic-kl (pre-merge outer, objective inner).
for pre, prelabel in (("maha", "Maha"), ("kl", "KL")):
    for obj in ("bic", "aic"):
        configs.append((f"{BASE}-smart-opt-{obj}-{pre}", f"{BASE}-smart-opt",
                        obj.upper(), f"optimal ({prelabel} pre-merge)", "hier"))

wb = openpyxl.load_workbook(XLSX)
ws = wb["Sheet1"]
tpl = {"hier": {c: ws.cell(4, c) for c in range(1, 13)}, "hdb": {c: ws.cell(34, c) for c in range(1, 13)}}


def style_like(dst, src):
    dst.font = copy.copy(src.font)
    dst.border = copy.copy(src.border)
    dst.fill = copy.copy(src.fill)
    dst.alignment = copy.copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy.copy(src.protection)


for r in range(MAHA_LAST_ROW + 1, ws.max_row + 3):  # unmerge any old footer rows, wherever they sit
    try:
        ws.unmerge_cells(f"A{r}:L{r}")
    except Exception:
        pass

filled = 0
row = MAHA_LAST_ROW + 1
for fname, version, c_label, d_method, tkey in configs:
    for c in range(1, 13):
        style_like(ws.cell(row, c), tpl[tkey][c])
    ws.cell(row, 1).value = f"=A{row - 1}+1"
    ws.cell(row, 2).value = version
    ws.cell(row, 3).value = c_label
    ws.cell(row, 4).value = d_method
    for col, b in ((6, "E"), (8, "G"), (10, "I"), (12, "K")):
        ws.cell(row, col).value = f'=IF(ISNUMBER({b}{row}),({b}{row}-{b}$3)/{b}$3,"")'
    d = load(fname)
    if d is not None and len(d) >= 1000:
        ll_u, rb_u, ll_r, rb_r = metrics(d)
        ws.cell(row, 5).value = round(ll_u, 6)
        ws.cell(row, 7).value = round(rb_u, 6)
        ws.cell(row, 9).value = round(ll_r, 6)
        ws.cell(row, 11).value = round(rb_r, 6)
        filled += 1
    else:
        for c in (5, 7, 9, 11):
            ws.cell(row, c).value = None
    row += 1

base_d = load(BASE)
total_reviews = sum(base_d[u]["size"] for u in base_d) if base_d else 0
ws.cell(row, 1).value = "1000 users"
ws.cell(row + 1, 1).value = f"{total_reviews:,} reviews (same-day reviews included)".replace(",", " ")
ws.merge_cells(f"A{row}:L{row}")
ws.merge_cells(f"A{row + 1}:L{row + 1}")

wb.save(XLSX)
print(f"appended {len(configs)} rows ({MAHA_LAST_ROW + 1}-{row - 1}), {filled} with data; footer {row}-{row + 1}")
