"""Fill Smart Preset Assignment.xlsx from the sweep results.

- Baseline (row 3) + 30 hierarchical (rows 4-33): write raw LL/RMSE into E/G/I/K; the Δ columns
  (F/H/J/L) are existing formulas that auto-compute relative% vs row 3.
- HDBSCAN: restructure the 6 eps rows into the 16 (mcs × ms × method) experiments (rows 34-49),
  shifting the footer notes to 50-51. Fills values if the hdbscan result files exist yet.
Idempotent: re-running rewrites the same cells (footer text is reconstructed, not read back).
"""
import json, os, copy, sys
import openpyxl

RDIR = sys.argv[1] if len(sys.argv) > 1 else "result"
BASE = "FSRS-7-short-secs"
XLSX = "Smart Preset Assignment.xlsx"

METHODS = ["single", "complete", "average", "centroid", "ward"]
THRESHOLDS = ["1.5", "2", "3", "5", "7.5", "12"]
HDB_MCS = [2, 5, 10, 20]
HDB_MS = [1, 5]

def load(name):
    p = f"{RDIR}/{name}.jsonl"
    if not os.path.exists(p):
        return None
    d = {}
    for line in open(p):
        line = line.strip()
        if not line:
            continue
        o = json.loads(line)
        if "metrics" in o and o["user"] <= 1000:
            d[o["user"]] = o
    return d or None

def metrics(d):
    """Return (LL_users, RMSE_users, LL_reviews, RMSE_reviews)."""
    us = sorted(d)
    ll = [d[u]["metrics"]["LogLoss"] for u in us]
    rb = [d[u]["metrics"]["RMSE(bins)"] for u in us]
    sz = [d[u]["size"] for u in us]
    n, tot = len(us), sum(sz)
    return (sum(ll) / n, sum(rb) / n,
            sum(v * s for v, s in zip(ll, sz)) / tot,
            sum(v * s for v, s in zip(rb, sz)) / tot)

wb = openpyxl.load_workbook(XLSX)
ws = wb["Sheet1"]

# Unmerge footer rows up front so rows 40/41 are writable by the HDBSCAN loop.
for mr in ("A40:L40", "A41:L41", "A50:L50", "A51:L51"):
    try:
        ws.unmerge_cells(mr)
    except Exception:
        pass

def set_vals(row, name):
    d = load(name)
    if d is None or len(d) < 1000:
        return False  # skip incomplete files (only fill once all 1000 users are present)
    ll_u, rb_u, ll_r, rb_r = metrics(d)
    ws.cell(row, 5).value = round(ll_u, 6)   # E log loss (users)
    ws.cell(row, 7).value = round(rb_u, 6)   # G RMSE bins (users)
    ws.cell(row, 9).value = round(ll_r, 6)   # I log loss (reviews)
    ws.cell(row, 11).value = round(rb_r, 6)  # K RMSE bins (reviews)
    return True

# 1. baseline + 30 hierarchical
set_vals(3, BASE)
for i, m in enumerate(METHODS):
    for j, t in enumerate(THRESHOLDS):
        set_vals(4 + i * 6 + j, f"{BASE}-smart-{m}-{t}")

# 2. HDBSCAN restructure: rows 34..49 (16 experiments), footer -> 50,51.
# Template styles from an existing hierarchical row (33).
tpl = {c: ws.cell(33, c) for c in range(1, 13)}
def style_like(dst, src):
    dst.font = copy.copy(src.font)
    dst.border = copy.copy(src.border)
    dst.fill = copy.copy(src.fill)
    dst.alignment = copy.copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy.copy(src.protection)

hdb = [(mcs, ms, meth) for mcs in HDB_MCS for ms in HDB_MS for meth in ("eom", "leaf")]
for k, (mcs, ms, meth) in enumerate(hdb):
    r = 34 + k
    for c in range(1, 13):
        style_like(ws.cell(r, c), tpl[c])
    ws.cell(r, 1).value = f"=A{r-1}+1"                       # A exp#
    ws.cell(r, 2).value = "FSRS-7-short-secs-smart"          # B version
    ws.cell(r, 3).value = f"mcs={mcs}, ms={ms}"              # C params
    ws.cell(r, 4).value = f"HDBSCAN ({meth})"               # D method
    for col, base_col in ((6, "E"), (8, "G"), (10, "I"), (12, "K")):  # F,H,J,L delta formulas
        ws.cell(r, col).value = f'=IF(ISNUMBER({base_col}{r}),({base_col}{r}-{base_col}$3)/{base_col}$3,"")'
    for c in (5, 7, 9, 11):  # clear value cells (filled below if results exist)
        ws.cell(r, c).value = None
    set_vals(r, f"{BASE}-smart-hdbscan-mcs{mcs}-ms{ms}-{meth}")

# footer (reconstructed so the script is idempotent)
base_d = load(BASE)
total_reviews = sum(base_d[u]["size"] for u in base_d) if base_d else 0
ws.cell(50, 1).value = "1000 users"
ws.cell(51, 1).value = f"{total_reviews:,} reviews (same-day reviews included)".replace(",", " ")
ws.merge_cells("A50:L50")
ws.merge_cells("A51:L51")

wb.save(XLSX)
print(f"filled xlsx; total reviews (1000 users) = {total_reviews:,}")
